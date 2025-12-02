"""Expense management and export routes."""

import base64
import io
from datetime import datetime

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

from api.schemas.expense import (
    ExtractedExpense,
    ExpenseValidationResult,
    ExportRequest,
    ExportResponse,
    ExpenseCategory,
)
from config import settings

router = APIRouter()


@router.post("/validate", response_model=ExpenseValidationResult)
async def validate_expense(expense: ExtractedExpense) -> ExpenseValidationResult:
    """
    Validate expense data for completeness and compliance.

    Checks:
    - Required fields are present
    - Amounts are valid
    - Hotel expenses have itemization
    - Meals over threshold have companion info
    """
    errors = []
    warnings = []

    # Required field validation
    if not expense.vendor or expense.vendor.strip() == "":
        errors.append("Vendor name is required")

    if expense.total <= 0:
        errors.append("Total amount must be greater than 0")

    # Date validation
    if expense.expense_date > datetime.now().date():
        warnings.append("Expense date is in the future")

    # Amount consistency
    expected_total = expense.subtotal + expense.tax
    if abs(float(expense.total) - float(expected_total)) > 0.01:
        warnings.append(
            f"Total ({expense.total}) doesn't match subtotal + tax ({expected_total})"
        )

    # Hotel itemization check
    if expense.category == ExpenseCategory.HOTEL:
        if not expense.hotel_itemization or len(expense.hotel_itemization) == 0:
            errors.append("Hotel expenses require per-night itemization")
        elif expense.check_in_date and expense.check_out_date:
            expected_nights = (expense.check_out_date - expense.check_in_date).days
            actual_nights = len(expense.hotel_itemization)
            if expected_nights != actual_nights:
                warnings.append(
                    f"Itemization has {actual_nights} nights but stay is {expected_nights} nights"
                )

    # Meal companion check
    if expense.category == ExpenseCategory.MEAL:
        if float(expense.total) > settings.meal_companion_threshold:
            if not expense.companion_info or not expense.companion_info.companion_name:
                errors.append(
                    f"Meals over ${settings.meal_companion_threshold} require companion information"
                )

    return ExpenseValidationResult(
        is_valid=len(errors) == 0,
        errors=errors,
        warnings=warnings,
    )


@router.post("/export", response_model=ExportResponse)
async def export_expenses(request: ExportRequest) -> ExportResponse:
    """
    Export expenses to Excel or CSV format.

    Generates a spreadsheet with:
    - Main summary sheet with all expenses
    - Hotel itemization sheet (if enabled)
    - Meals companion sheet (if enabled)
    """
    try:
        if request.format == "xlsx":
            content, filename = generate_excel_export(
                request.expenses,
                request.include_itemization,
                request.include_companion_sheet,
            )
        else:
            content, filename = generate_csv_export(request.expenses)

        content_base64 = base64.b64encode(content).decode("utf-8")

        return ExportResponse(
            success=True,
            file_name=filename,
            file_content_base64=content_base64,
            error_message=None,
        )

    except Exception as e:
        return ExportResponse(
            success=False,
            file_name="",
            file_content_base64="",
            error_message=f"Export failed: {str(e)}",
        )


@router.post("/export/download")
async def download_export(request: ExportRequest):
    """
    Export and directly download expenses as a file.

    Returns the file as a streaming response for direct download.
    """
    try:
        if request.format == "xlsx":
            content, filename = generate_excel_export(
                request.expenses,
                request.include_itemization,
                request.include_companion_sheet,
            )
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            content, filename = generate_csv_export(request.expenses)
            media_type = "text/csv"

        return StreamingResponse(
            io.BytesIO(content),
            media_type=media_type,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Export failed: {str(e)}")


def generate_excel_export(
    expenses: list[ExtractedExpense],
    include_itemization: bool,
    include_companion: bool,
) -> tuple[bytes, str]:
    """Generate Excel export with multiple sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()

    # Main summary sheet
    ws = wb.active
    ws.title = "Expense Summary"

    # Headers
    headers = [
        "Date",
        "Category",
        "Vendor",
        "Description",
        "Subtotal",
        "Tax",
        "Total",
        "Currency",
        "Payment Method",
        "Receipt #",
        "Business Purpose",
        "Project Code",
    ]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row, expense in enumerate(expenses, 2):
        ws.cell(row=row, column=1, value=str(expense.expense_date))
        ws.cell(row=row, column=2, value=expense.category.value)
        ws.cell(row=row, column=3, value=expense.vendor)
        ws.cell(row=row, column=4, value=expense.description or "")
        ws.cell(row=row, column=5, value=float(expense.subtotal))
        ws.cell(row=row, column=6, value=float(expense.tax))
        ws.cell(row=row, column=7, value=float(expense.total))
        ws.cell(row=row, column=8, value=expense.currency)
        ws.cell(row=row, column=9, value=expense.payment_method or "")
        ws.cell(row=row, column=10, value=expense.receipt_number or "")
        ws.cell(row=row, column=11, value=expense.business_purpose or "")
        ws.cell(row=row, column=12, value=expense.project_code or "")

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    # Hotel itemization sheet
    if include_itemization:
        hotel_expenses = [e for e in expenses if e.category == ExpenseCategory.HOTEL]
        if hotel_expenses:
            ws_hotel = wb.create_sheet("Hotel Itemization")

            hotel_headers = [
                "Hotel",
                "Check-in",
                "Check-out",
                "Night Date",
                "Room Rate",
                "Room Tax",
                "Service Charge",
                "Resort Fee",
                "Parking",
                "Other Fees",
                "Daily Total",
            ]

            for col, header in enumerate(hotel_headers, 1):
                cell = ws_hotel.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            row = 2
            for expense in hotel_expenses:
                if expense.hotel_itemization:
                    for i, night in enumerate(expense.hotel_itemization):
                        ws_hotel.cell(
                            row=row, column=1, value=expense.vendor if i == 0 else ""
                        )
                        ws_hotel.cell(
                            row=row,
                            column=2,
                            value=str(expense.check_in_date) if i == 0 else "",
                        )
                        ws_hotel.cell(
                            row=row,
                            column=3,
                            value=str(expense.check_out_date) if i == 0 else "",
                        )
                        ws_hotel.cell(row=row, column=4, value=str(night.night_date))
                        ws_hotel.cell(row=row, column=5, value=float(night.room_rate))
                        ws_hotel.cell(row=row, column=6, value=float(night.room_tax))
                        ws_hotel.cell(row=row, column=7, value=float(night.service_charge))
                        ws_hotel.cell(row=row, column=8, value=float(night.resort_fee))
                        ws_hotel.cell(row=row, column=9, value=float(night.parking_fee))
                        ws_hotel.cell(row=row, column=10, value=float(night.other_fees))
                        ws_hotel.cell(row=row, column=11, value=float(night.daily_total))
                        row += 1
                    row += 1  # Empty row between hotels

    # Meals companion sheet
    if include_companion:
        meal_expenses = [
            e
            for e in expenses
            if e.category == ExpenseCategory.MEAL
            and float(e.total) > settings.meal_companion_threshold
        ]
        if meal_expenses:
            ws_meals = wb.create_sheet("Meals Entertainment")

            meal_headers = [
                "Date",
                "Vendor",
                "Total",
                "Companion Name",
                "Companion Title",
                "Attendees",
                "Business Purpose",
                "Discussion Topics",
            ]

            for col, header in enumerate(meal_headers, 1):
                cell = ws_meals.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font

            for row, expense in enumerate(meal_expenses, 2):
                ws_meals.cell(row=row, column=1, value=str(expense.expense_date))
                ws_meals.cell(row=row, column=2, value=expense.vendor)
                ws_meals.cell(row=row, column=3, value=float(expense.total))

                if expense.companion_info:
                    ws_meals.cell(
                        row=row, column=4, value=expense.companion_info.companion_name
                    )
                    ws_meals.cell(
                        row=row, column=5, value=expense.companion_info.companion_title or ""
                    )
                    ws_meals.cell(
                        row=row, column=6, value=expense.companion_info.attendee_count
                    )
                    ws_meals.cell(
                        row=row, column=7, value=expense.companion_info.business_purpose
                    )
                    ws_meals.cell(
                        row=row,
                        column=8,
                        value=expense.companion_info.discussion_topics or "",
                    )
                else:
                    ws_meals.cell(row=row, column=4, value="NOT PROVIDED")

    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"expense_report_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

    return output.getvalue(), filename


def generate_csv_export(expenses: list[ExtractedExpense]) -> tuple[bytes, str]:
    """Generate CSV export."""
    import csv

    output = io.StringIO()
    writer = csv.writer(output)

    # Headers
    headers = [
        "Date",
        "Category",
        "Vendor",
        "Description",
        "Subtotal",
        "Tax",
        "Total",
        "Currency",
        "Payment Method",
        "Receipt #",
        "Business Purpose",
        "Project Code",
    ]
    writer.writerow(headers)

    # Data rows
    for expense in expenses:
        writer.writerow(
            [
                str(expense.expense_date),
                expense.category.value,
                expense.vendor,
                expense.description or "",
                float(expense.subtotal),
                float(expense.tax),
                float(expense.total),
                expense.currency,
                expense.payment_method or "",
                expense.receipt_number or "",
                expense.business_purpose or "",
                expense.project_code or "",
            ]
        )

    filename = f"expense_report_{datetime.now().strftime('%Y-%m-%d')}.csv"

    return output.getvalue().encode("utf-8"), filename


@router.get("/categories")
async def get_categories() -> dict:
    """Get list of available expense categories."""
    return {
        "categories": [
            {"value": cat.value, "label": cat.value.replace("_", " ").title()}
            for cat in ExpenseCategory
        ]
    }


@router.get("/business-rules")
async def get_business_rules() -> dict:
    """Get business rules for expense processing."""
    return {
        "meal_companion_threshold": settings.meal_companion_threshold,
        "requires_itemization": ["hotel"],
        "requires_companion_over_threshold": ["meal"],
    }
